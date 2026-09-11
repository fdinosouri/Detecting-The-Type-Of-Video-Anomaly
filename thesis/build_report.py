# -*- coding: utf-8 -*-
"""Assemble the four-chapter project report into a Word document.

    python thesis/make_figures.py    # once, to draw the plots
    python thesis/build_report.py

Layout follows the same Hamedan University of Technology writing guide the
thesis uses: A4 with the prescribed margins, front matter without page
numbers, the contents pages numbered with abjad letters, the main text
numbered with digits, running headers per chapter, numbered captions above
tables and below figures, and equations numbered per chapter.

The report differs from the thesis only in what it contains: four chapters
instead of five, no appendices, and no dedication page. Styles, metadata,
acknowledgement and references come from the same modules, so the two
documents cannot drift apart.
"""

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build_thesis as B
import content_report_fa as RC
from docx_style import add_section, set_footer_page_number, set_header_text

OUTPUT = HERE / "Report_HUT_VideoAnomaly.docx"


class ReportBuilder(B.Builder):
    """The thesis builder with a report's front and back matter."""

    def title_page_fa(self):
        B.add_page_break(self.doc)
        meta = RC.META

        self.para(meta["university"], style="Thesis Title Line",
                  sizes=(16, 15), bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para(meta["department"], style="Thesis Title Line",
                  sizes=(14, 13), align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank(1)
        self.para(
            f"گزارش کار پروژه‌ی {meta['degree']} در رشته‌ی {meta['field']} – "
            f"گرایش {meta['orientation']}",
            sizes=(13, 12), align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank(1)
        self.para(meta["title"], style="Thesis Title Line", sizes=(18, 16),
                  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank(1)
        self.para("توسط:", sizes=(14, 13), align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para(meta["author"], style="Thesis Title Line", sizes=(16, 15),
                  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank(2)
        self.para("استاد راهنما:", sizes=(14, 13),
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para(meta["supervisor"], style="Thesis Title Line",
                  sizes=(16, 15), bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank(2)
        self.para(meta["date"], sizes=(14, 13),
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    def abstract_en(self):
        section = add_section(self.doc, rtl=False, different_first_page=False)
        set_header_text(section, "")
        set_footer_page_number(section, enabled=False)

        self.para("Abstract", style="Thesis Title Line EN", sizes=(18, 16),
                  bold=True, rtl=False, align=WD_ALIGN_PARAGRAPH.CENTER)

        for paragraph in RC.ABSTRACT_EN:
            self.para(paragraph, style="Thesis Abstract EN", rtl=False,
                      indent=True)

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = B.Pt(12)
        spacer.paragraph_format.line_spacing = 1.0

        paragraph = self.doc.add_paragraph(style="Thesis Abstract EN")
        B.ltr_paragraph(paragraph)
        label = paragraph.add_run("Keywords: ")
        B.set_run_font(label, B.SIZE_BODY, bold=True, rtl=False)
        body = paragraph.add_run(RC.KEYWORDS_EN)
        B.set_run_font(body, B.SIZE_BODY, rtl=False)

    def build(self):
        set_footer_page_number(self.doc.sections[0], enabled=False)

        self.bismillah()
        self.title_page_fa()
        self.acknowledgement()
        self.abstract_fa()

        self.contents()

        for index, chapter in enumerate(RC.CHAPTERS, start=1):
            self.chapter(index, chapter, first=(index == 1))

        self.references()
        self.abstract_en()
        self.title_page_en()

        self.doc.save(OUTPUT)

        return OUTPUT


def main():
    # The inherited front matter, contents and references read their text
    # from build_thesis's content module, so point it at the report's.
    B.C = RC

    path = ReportBuilder().build()
    size = path.stat().st_size / 1024
    print(f"wrote {path}  ({size:.0f} KB)")
    print("open it in Word, press Ctrl+A then F9 to fill the three contents "
          "pages, and install the B Nazanin font if it is not present.")


if __name__ == "__main__":
    main()
